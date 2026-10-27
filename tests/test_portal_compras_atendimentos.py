import sqlite3
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

import database.connection as dbc
from modules.requisicoes_compra import services as svc
from modules.requisicoes_compra.atendimentos import estornar_atendimento, gerar_excel, registrar_atendimento, relatorio_atendidas, relatorio_recusadas


@pytest.fixture()
def banco(tmp_path,monkeypatch):
    path=str(tmp_path/"portal.sqlite3"); monkeypatch.setattr(dbc,"DB_NAME",path)
    conn=sqlite3.connect(path); conn.executescript("""
    CREATE TABLE almoxarifado_insumos(id INTEGER PRIMARY KEY,descricao TEXT,categoria TEXT,unidade TEXT,ativo TEXT,origem_baixa TEXT);
    CREATE TABLE parceiros(id INTEGER PRIMARY KEY,razao_social TEXT,nome_fantasia TEXT,status TEXT);
    CREATE TABLE parceiro_papeis(id INTEGER PRIMARY KEY,parceiro_id INTEGER,papel TEXT,ativo INTEGER);
    CREATE TABLE almoxarifado_movimentacoes(id INTEGER PRIMARY KEY,quantidade REAL);
    CREATE TABLE movimentacoes_financeiras(id INTEGER PRIMARY KEY,valor REAL);
    INSERT INTO almoxarifado_insumos VALUES(1,'Caixa','Embalagem','un','Sim','ORDEM_PRODUCAO');
    INSERT INTO parceiros VALUES(1,'Fornecedor A','Forn A','Ativo'),(2,'Fornecedor B',NULL,'Ativo');
    INSERT INTO parceiro_papeis VALUES(1,1,'FORNECEDOR',1),(2,2,'FORNECEDOR',1);
    """); conn.commit(); conn.close(); svc.criar_tabelas_requisicoes_compra(); return path

def u(i=1,perfil="admin"): return {"id":i,"nome":f"U{i}","perfil":perfil}
def criar_aprovada(chave="rc",qtd="10",custo="10"):
    dados={"tipo_origem":"ADMINISTRATIVO","setor":"Administrativo","origem_descricao":"Compra","prioridade":"NORMAL","justificativa":"Necessidade operacional"}; item={"material_id":"1","quantidade_item":qtd,"custo_item":custo,"descricao_item":"","unidade_item":"","observacao_item":""}
    rc=svc.criar_rascunho(dados,[item],ator=u(),idempotency_key=chave); rc=svc.enviar(rc["id"],ator=u(),versao=0,idempotency_key=chave+"-env"); return svc.aprovar(rc["id"],ator=u(2,"gerencia"),versao=1,idempotency_key=chave+"-apr")
def compra(rc,qtd,preco,chave,fornecedor=1):
    d={"tipo_documento":"NOTA_FISCAL","numero_documento":chave,"serie":"1","data_emissao":"2026-10-27","data_atendimento":"2026-10-27","fornecedor_id":str(fornecedor),"valor_total_documento":str(Decimal(qtd)*Decimal(preco)),"observacao":"Compra testada"}
    return registrar_atendimento(rc["id"],d,[{"item_id":rc["itens"][0]["id"],"quantidade":qtd,"preco_unitario_real":preco}],ator=u(3,"pcp"),idempotency_key=chave)

def test_parcial_segundo_documento_completa_estimado_preservado_e_sem_efeitos_colaterais(banco):
    rc=criar_aprovada(); estimado=rc["total_estimado_aprovado"]
    conn=sqlite3.connect(banco); antes=(conn.execute("SELECT COUNT(*) FROM almoxarifado_movimentacoes").fetchone()[0],conn.execute("SELECT COUNT(*) FROM movimentacoes_financeiras").fetchone()[0]); conn.close()
    rc=compra(rc,"6","9","nf1"); assert rc["status"]=="PARCIALMENTE_ATENDIDA" and rc["itens"][0]["quantidade_pendente"]==Decimal("4.0")
    rc=compra(rc,"4","12","nf2",2); assert rc["status"]=="ATENDIDA" and len(rc["atendimentos"])==2 and rc["total_estimado_aprovado"]==estimado and rc["custo_real_acumulado"]==Decimal("102.0")
    conn=sqlite3.connect(banco); depois=(conn.execute("SELECT COUNT(*) FROM almoxarifado_movimentacoes").fetchone()[0],conn.execute("SELECT COUNT(*) FROM movimentacoes_financeiras").fetchone()[0]); conn.close(); assert depois==antes

def test_limite_idempotencia_permissao_e_documento_obrigatorio(banco):
    rc=criar_aprovada("limite"); compra(rc,"6","10","doc-a")
    with pytest.raises(ValueError,match="excede"): compra(rc,"5","10","doc-b")
    repetido=compra(rc,"6","10","doc-a"); assert len(repetido["atendimentos"])==1
    with pytest.raises(PermissionError): registrar_atendimento(rc["id"],{},[],ator=u(9,"producao"),idempotency_key="negado")
    with pytest.raises(ValueError,match="obrigatórios"): registrar_atendimento(rc["id"],{"tipo_documento":"NOTA_FISCAL"},[],ator=u(3,"pcp"),idempotency_key="sem-doc")

def test_estorno_formal_reabre_saldo_e_audita(banco):
    rc=compra(criar_aprovada("estorno"),"10","10","nf-est"); atendimento=rc["atendimentos"][0]
    with pytest.raises(PermissionError): estornar_atendimento(rc["id"],atendimento["id"],"Erro",ator=u(3,"pcp"),idempotency_key="est-neg")
    rc=estornar_atendimento(rc["id"],atendimento["id"],"Documento incorreto",ator=u(2,"gerencia"),idempotency_key="est-ok")
    assert rc["status"]=="APROVADA" and rc["atendimentos"][0]["status"]=="ESTORNADO" and any(e["evento"]=="ESTORNO_ATENDIMENTO" for e in rc["eventos"])

def test_relatorios_calculos_e_excel(banco):
    compra(criar_aprovada("economia"),"10","9","nf-econ")
    compra(criar_aprovada("estouro"),"10","12","nf-estouro",2)
    zero=criar_aprovada("zero",qtd="1",custo="0"); compra(zero,"1","5","nf-zero")
    dados=relatorio_atendidas({}); assert dados["indicadores"]["economia"]==Decimal("10.00") and dados["indicadores"]["estouro"]==Decimal("25.00")
    assert next(x for x in dados["linhas"] if x["numero"]==zero["numero"])["variacao"] is None
    wb=load_workbook(BytesIO(gerar_excel("atendidas",{}))); assert wb.active["A1"].value=="RC" and wb.active.max_row==4

def test_recusadas_gasto_evitado_e_filtros(banco):
    rc=criar_aprovada("base-rej")
    # Uma RC separada é recusada antes da aprovação.
    dados={"tipo_origem":"ADMINISTRATIVO","setor":"TI","origem_descricao":"Compra","prioridade":"NORMAL","justificativa":"Necessidade operacional"}; item={"material_id":"1","quantidade_item":"2","custo_item":"25","descricao_item":"","unidade_item":"","observacao_item":""}
    rejeitada=svc.criar_rascunho(dados,[item],ator=u(),idempotency_key="rej"); rejeitada=svc.enviar(rejeitada["id"],ator=u(),versao=0,idempotency_key="rej-env"); svc.rejeitar(rejeitada["id"],ator=u(2,"gerencia"),versao=1,idempotency_key="rej-fim",motivo="Sem orçamento")
    rel=relatorio_recusadas({"setor":"TI","motivo":"orçamento"}); assert rel["indicadores"]["quantidade"]==1 and rel["indicadores"]["total"]==Decimal("50.00") and rel["indicadores"]["setor"]=="TI"
    assert load_workbook(BytesIO(gerar_excel("recusadas",{}))).active.max_row==2 and rc["status"]=="APROVADA"

def test_migration_sqlite_upgrade_downgrade_upgrade(banco):
    root=Path(__file__).resolve().parents[1]/"database"; up=(root/"20261027_portal_compras_atendimentos_sqlite.sql").read_text(); down=(root/"20261027_portal_compras_atendimentos_sqlite_rollback.sql").read_text(); conn=sqlite3.connect(banco)
    conn.executescript(up); conn.executescript(up); conn.executescript(down); conn.executescript(up); tabelas={x[0] for x in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}; conn.close(); assert {"documentos_compra","rc_atendimentos","rc_atendimento_itens"}<=tabelas
