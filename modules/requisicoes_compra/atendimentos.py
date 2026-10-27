"""Atendimento documental de RC, sem efeitos em estoque ou Financeiro."""
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from database import DATABASE_URL, conectar, q

CENTAVOS=Decimal("0.01")
TIPOS_DOCUMENTO=("NOTA_FISCAL","RECIBO","CUPOM","FATURA","OUTRO")
STATUS_COMPRA=("APROVADA","APROVADA_COM_AJUSTES","PARCIALMENTE_ATENDIDA")
PERFIS_ATENDIMENTO={"admin","pcp","gerencia"}
PERFIS_ESTORNO={"admin","gerencia"}

def _agora(): return datetime.now().replace(microsecond=0).isoformat(sep=" ")
def _u(ator): return {"id":ator.get("id") or ator.get("usuario_id"),"nome":str(ator.get("nome") or "Sistema"),"perfil":str(ator.get("perfil") or "").lower()}
def _dec(v,campo,zero=False):
    try: n=Decimal(str(v or "").replace(",","."))
    except (InvalidOperation,ValueError): raise ValueError(f"{campo} inválido.") from None
    if not n.is_finite() or n<0 or (not zero and n==0): raise ValueError(f"{campo} deve ser maior que zero.")
    return n
def _id(cur):
    if DATABASE_URL: cur.execute("SELECT LASTVAL() id"); return int(cur.fetchone()["id"])
    return int(cur.lastrowid)
def _begin(conn):
    if not DATABASE_URL: conn.execute("BEGIN IMMEDIATE")
def _fornecedor(cur,fornecedor_id):
    cur.execute(q("""SELECT p.id,p.razao_social,p.nome_fantasia FROM parceiros p JOIN parceiro_papeis pp ON pp.parceiro_id=p.id AND pp.papel='FORNECEDOR' AND pp.ativo=1 WHERE p.id=? AND p.status='Ativo'"""),(fornecedor_id,)); row=cur.fetchone()
    if not row: raise ValueError("Fornecedor ativo não encontrado no Cadastro de Parceiros.")
    return dict(row)
def listar_fornecedores():
    conn=conectar()
    try:
        cur=conn.cursor(); cur.execute("""SELECT DISTINCT p.id,p.razao_social,p.nome_fantasia FROM parceiros p JOIN parceiro_papeis pp ON pp.parceiro_id=p.id AND pp.papel='FORNECEDOR' AND pp.ativo=1 WHERE p.status='Ativo' ORDER BY p.razao_social"""); return [dict(x) for x in cur.fetchall()]
    finally: conn.close()

def registrar_atendimento(rid,dados,itens,*,ator,idempotency_key):
    u=_u(ator); chave=str(idempotency_key or "").strip()
    if u["perfil"] not in PERFIS_ATENDIMENTO: raise PermissionError("Sem permissão para registrar atendimento.")
    if not chave: raise ValueError("Chave idempotente obrigatória.")
    tipo=str(dados.get("tipo_documento") or "").upper(); numero=str(dados.get("numero_documento") or "").strip(); emissao=str(dados.get("data_emissao") or "").strip(); data=str(dados.get("data_atendimento") or emissao).strip()
    if tipo not in TIPOS_DOCUMENTO or not numero or not emissao or not data: raise ValueError("Tipo, número e datas do documento são obrigatórios.")
    total_doc=_dec(dados.get("valor_total_documento"),"Valor total do documento",True); fornecedor_id=int(dados.get("fornecedor_id") or 0)
    preparados=[]
    for raw in itens:
        qtd=str(raw.get("quantidade") or "").strip(); preco=str(raw.get("preco_unitario_real") or "").strip()
        if not qtd and not preco: continue
        preparados.append((int(raw.get("item_id") or 0),_dec(qtd,"Quantidade comprada"),_dec(preco,"Preço unitário real",True)))
    if not preparados: raise ValueError("Informe ao menos um item atendido.")
    conn=conectar()
    try:
        _begin(conn); cur=conn.cursor(); cur.execute(q("SELECT id FROM rc_atendimentos WHERE chave_criacao=?"),(chave,)); repetido=cur.fetchone()
        if repetido: conn.rollback(); from .services import buscar_rc; return buscar_rc(rid)
        lock=" FOR UPDATE" if DATABASE_URL else ""; cur.execute(q(f"SELECT * FROM requisicoes_compra WHERE id=?{lock}"),(rid,)); rc=cur.fetchone()
        if not rc or rc["status"] not in STATUS_COMPRA: raise ValueError("RC não está disponível para atendimento.")
        fornecedor=_fornecedor(cur,fornecedor_id); ids=[x[0] for x in preparados]; marks=",".join("?" for _ in ids)
        cur.execute(q(f"SELECT * FROM requisicao_compra_itens WHERE requisicao_compra_id=? AND id IN ({marks}){lock}"),tuple([rid,*ids])); linhas={int(x["id"]):dict(x) for x in cur.fetchall()}
        if len(linhas)!=len(set(ids)): raise ValueError("Item de atendimento inválido.")
        t=_agora(); nome=fornecedor.get("nome_fantasia") or fornecedor["razao_social"]
        cur.execute(q("""INSERT INTO documentos_compra(requisicao_compra_id,tipo,numero,serie,data_emissao,fornecedor_id,fornecedor_nome_snapshot,valor_total,observacao,status,criado_por,criado_por_nome,criado_em,chave_criacao) VALUES(?,?,?,?,?,?,?,?,?,'ATIVO',?,?,?,?)"""),(rid,tipo,numero,str(dados.get("serie") or "").strip() or None,emissao,fornecedor_id,nome,float(total_doc),str(dados.get("observacao") or "").strip() or None,u["id"],u["nome"],t,chave+":doc")); documento_id=_id(cur)
        cur.execute(q("""INSERT INTO rc_atendimentos(requisicao_compra_id,documento_id,data_atendimento,observacao,status,usuario_id,usuario_nome,criado_em,chave_criacao) VALUES(?,?,?,?,'ATIVO',?,?,?,?)"""),(rid,documento_id,data,str(dados.get("observacao") or "").strip() or None,u["id"],u["nome"],t,chave)); atendimento_id=_id(cur)
        valores=[]
        for item_id,qtd,preco in preparados:
            item=linhas[item_id]; cur.execute(q("""SELECT COALESCE(SUM(ai.quantidade),0) total FROM rc_atendimento_itens ai JOIN rc_atendimentos a ON a.id=ai.atendimento_id WHERE ai.requisicao_compra_item_id=? AND a.status='ATIVO'"""),(item_id,)); atendido=Decimal(str(cur.fetchone()["total"] or 0)); aprovado=Decimal(str(item["quantidade_aprovada"] if item["quantidade_aprovada"] is not None else item["quantidade_solicitada"]))
            if atendido+qtd>aprovado: raise ValueError(f"Quantidade de {item['descricao_snapshot']} excede o saldo aprovado de {aprovado-atendido}.")
            valor=(qtd*preco).quantize(CENTAVOS,rounding=ROUND_HALF_UP); cur.execute(q("INSERT INTO rc_atendimento_itens(atendimento_id,requisicao_compra_item_id,quantidade,preco_unitario_real,valor_real) VALUES(?,?,?,?,?)"),(atendimento_id,item_id,float(qtd),float(preco),float(valor))); valores.append({"item_id":item_id,"quantidade":str(qtd),"preco":str(preco),"valor":str(valor)})
        novo=_recalcular_status(cur,rid)
        cur.execute(q("UPDATE requisicoes_compra SET status=?,versao=versao+1,atualizado_em=? WHERE id=?"),(novo,t,rid))
        cur.execute(q("""INSERT INTO requisicao_compra_eventos(requisicao_compra_id,evento,status_anterior,status_novo,dados_novos,usuario_id,usuario_nome,perfil,criado_em,idempotency_key) VALUES(?,?,?,?,?,?,?,?,?,?)"""),(rid,"ATENDIMENTO_COMPRA",rc["status"],novo,json.dumps({"atendimento_id":atendimento_id,"documento_id":documento_id,"itens":valores},ensure_ascii=False),u["id"],u["nome"],u["perfil"],t,chave+":evento")); conn.commit()
        from .services import buscar_rc; return buscar_rc(rid)
    except Exception: conn.rollback(); raise
    finally: conn.close()

def _recalcular_status(cur,rid):
    cur.execute(q("""SELECT i.id,i.quantidade_solicitada,i.quantidade_aprovada,COALESCE(SUM(CASE WHEN a.status='ATIVO' THEN ai.quantidade ELSE 0 END),0) atendido FROM requisicao_compra_itens i LEFT JOIN rc_atendimento_itens ai ON ai.requisicao_compra_item_id=i.id LEFT JOIN rc_atendimentos a ON a.id=ai.atendimento_id WHERE i.requisicao_compra_id=? GROUP BY i.id,i.quantidade_solicitada,i.quantidade_aprovada"""),(rid,)); linhas=cur.fetchall()
    if linhas and all(Decimal(str(x["atendido"]))==Decimal(str(x["quantidade_aprovada"] if x["quantidade_aprovada"] is not None else x["quantidade_solicitada"])) for x in linhas): return "ATENDIDA"
    if any(Decimal(str(x["atendido"]))>0 for x in linhas): return "PARCIALMENTE_ATENDIDA"
    return "APROVADA_COM_AJUSTES" if any(x["quantidade_aprovada"] is not None and Decimal(str(x["quantidade_aprovada"]))!=Decimal(str(x["quantidade_solicitada"])) for x in linhas) else "APROVADA"

def estornar_atendimento(rid,atendimento_id,motivo,*,ator,idempotency_key):
    u=_u(ator); motivo=str(motivo or "").strip(); chave=str(idempotency_key or "").strip()
    if u["perfil"] not in PERFIS_ESTORNO: raise PermissionError("Somente administração ou gerência pode estornar.")
    if not motivo or not chave: raise ValueError("Motivo e chave idempotente são obrigatórios.")
    conn=conectar()
    try:
        _begin(conn); cur=conn.cursor(); lock=" FOR UPDATE" if DATABASE_URL else ""; cur.execute(q(f"SELECT * FROM rc_atendimentos WHERE id=? AND requisicao_compra_id=?{lock}"),(atendimento_id,rid)); a=cur.fetchone()
        cur.execute(q("SELECT id FROM requisicao_compra_eventos WHERE idempotency_key=?"),(chave,));
        if cur.fetchone(): conn.rollback(); from .services import buscar_rc; return buscar_rc(rid)
        if not a or a["status"]!="ATIVO": raise ValueError("Atendimento não está ativo.")
        t=_agora(); cur.execute(q("UPDATE rc_atendimentos SET status='ESTORNADO',estornado_por=?,estornado_por_nome=?,estornado_em=?,motivo_estorno=? WHERE id=?"),(u["id"],u["nome"],t,motivo,atendimento_id)); cur.execute(q("UPDATE documentos_compra SET status='ESTORNADO',estornado_por=?,estornado_por_nome=?,estornado_em=?,motivo_estorno=? WHERE id=?"),(u["id"],u["nome"],t,motivo,a["documento_id"])); novo=_recalcular_status(cur,rid); cur.execute(q("SELECT status FROM requisicoes_compra WHERE id=?"),(rid,)); ant=cur.fetchone()["status"]; cur.execute(q("UPDATE requisicoes_compra SET status=?,versao=versao+1,atualizado_em=? WHERE id=?"),(novo,t,rid)); cur.execute(q("""INSERT INTO requisicao_compra_eventos(requisicao_compra_id,evento,status_anterior,status_novo,justificativa,usuario_id,usuario_nome,perfil,criado_em,idempotency_key) VALUES(?,'ESTORNO_ATENDIMENTO',?,?,?,?,?,?,?,?)"""),(rid,ant,novo,motivo,u["id"],u["nome"],u["perfil"],t,chave)); conn.commit(); from .services import buscar_rc; return buscar_rc(rid)
    except Exception: conn.rollback(); raise
    finally: conn.close()

def enriquecer_rc(cur,rc):
    rid=rc["id"]; cur.execute(q("""SELECT a.*,d.tipo,d.numero,d.serie,d.data_emissao,d.fornecedor_id,d.fornecedor_nome_snapshot,d.valor_total,d.status documento_status,ai.id atendimento_item_id,ai.requisicao_compra_item_id,ai.quantidade,ai.preco_unitario_real,ai.valor_real FROM rc_atendimentos a JOIN documentos_compra d ON d.id=a.documento_id LEFT JOIN rc_atendimento_itens ai ON ai.atendimento_id=a.id WHERE a.requisicao_compra_id=? ORDER BY a.id,ai.id"""),(rid,)); linhas=[dict(x) for x in cur.fetchall()]; rc["atendimentos"]=[]; vistos=set(); vinculados=[]
    for x in linhas:
        if x["id"] not in vistos: rc["atendimentos"].append({k:v for k,v in x.items() if k not in {"atendimento_item_id","requisicao_compra_item_id","quantidade","preco_unitario_real","valor_real"}}); vistos.add(x["id"])
        if x.get("atendimento_item_id") is not None: vinculados.append({"atendimento_id":x["id"],"requisicao_compra_item_id":x["requisicao_compra_item_id"],"quantidade":x["quantidade"],"preco_unitario_real":x["preco_unitario_real"],"valor_real":x["valor_real"],"atendimento_status":x["status"]})
    rc["custo_real_acumulado"]=Decimal("0")
    for item in rc["itens"]:
        ativos=[x for x in vinculados if x["requisicao_compra_item_id"]==item["id"] and x["atendimento_status"]=="ATIVO"]; item["quantidade_atendida"]=sum(Decimal(str(x["quantidade"])) for x in ativos); item["quantidade_pendente"]=Decimal(str(item["quantidade_efetiva"]))-item["quantidade_atendida"]; item["valor_real"]=sum(Decimal(str(x["valor_real"])) for x in ativos); rc["custo_real_acumulado"]+=item["valor_real"]; item["preco_real_medio"]=(item["valor_real"]/item["quantidade_atendida"]) if item["quantidade_atendida"] else None; item["diferenca_real"]=item["valor_real"]-(item["total_estimado_aprovado"] or Decimal("0"))
    rc["saldo_estimado_pendente"]=sum((Decimal(str(i["quantidade_pendente"]))*Decimal(str(i["custo_estimado_unitario"] or 0)) for i in rc["itens"]),Decimal("0"))
    nomes={i["id"]:i["descricao_snapshot"] for i in rc["itens"]}
    for atendimento in rc["atendimentos"]:
        atendimento["itens_resumo"]="; ".join(f"{nomes.get(x['requisicao_compra_item_id'],'Item')}: {x['quantidade']}" for x in vinculados if x["atendimento_id"]==atendimento["id"])

def _filtros_relatorio(f):
    cond=["1=1"]; p=[]
    for key,col in (("data_inicio","r.criado_em >="),("data_fim","r.criado_em <="),("setor","r.setor ="),("solicitante","r.solicitante_nome_snapshot =")):
        v=str(f.get(key) or "").strip()
        if v: cond.append(f"{col} ?"); p.append(v+(" 23:59:59" if key=="data_fim" else ""))
    return cond,p

def relatorio_atendidas(filtros=None):
    f=filtros or {}; cond,p=_filtros_relatorio(f); cond.append("r.status IN ('PARCIALMENTE_ATENDIDA','ATENDIDA')")
    if f.get("status"): cond.append("r.status=?"); p.append(f["status"])
    if f.get("fornecedor_id"): cond.append("EXISTS(SELECT 1 FROM documentos_compra d WHERE d.requisicao_compra_id=r.id AND d.status='ATIVO' AND d.fornecedor_id=?)"); p.append(f["fornecedor_id"])
    if f.get("item"): cond.append("EXISTS(SELECT 1 FROM requisicao_compra_itens i WHERE i.requisicao_compra_id=r.id AND LOWER(i.descricao_snapshot) LIKE ?)"); p.append("%"+str(f["item"]).lower()+"%")
    conn=conectar()
    try:
        cur=conn.cursor(); cur.execute(q(f"""SELECT r.* FROM requisicoes_compra r WHERE {' AND '.join(cond)} ORDER BY r.id DESC"""),tuple(p)); linhas=[]
        for row in cur.fetchall():
            rc=dict(row); cur.execute(q("SELECT * FROM requisicao_compra_itens WHERE requisicao_compra_id=?"),(rc["id"],)); rc["itens"]=[dict(x) for x in cur.fetchall()]; from .services import _totais; _totais(rc); enriquecer_rc(cur,rc); ativos=[a for a in rc["atendimentos"] if a["status"]=="ATIVO"]; estimado=rc["total_estimado_aprovado"]; real=rc["custo_real_acumulado"]; dif=real-estimado; variacao=None if estimado==0 else dif/estimado*100; linhas.append({**rc,"valor_estimado":estimado,"valor_real":real,"diferenca":dif,"variacao":variacao,"fornecedores":", ".join(sorted({a["fornecedor_nome_snapshot"] for a in ativos})),"documentos":", ".join(f"{a['tipo']} {a['numero']}" for a in ativos),"data_atendimento":max((a["data_atendimento"] for a in ativos),default=None)})
        economia=sum((-x["diferenca"] for x in linhas if x["diferenca"]<0),Decimal("0")); estouro=sum((x["diferenca"] for x in linhas if x["diferenca"]>0),Decimal("0")); calculaveis=[x["variacao"] for x in linhas if x["variacao"] is not None]; return {"linhas":linhas,"indicadores":{"quantidade":sum(1 for x in linhas if x["status"]=="ATENDIDA"),"estimado":sum((x["valor_estimado"] for x in linhas),Decimal("0")),"real":sum((x["valor_real"] for x in linhas),Decimal("0")),"economia":economia,"estouro":estouro,"variacao_media":sum(calculaveis,Decimal("0"))/len(calculaveis) if calculaveis else None}}
    finally: conn.close()

def relatorio_recusadas(filtros=None):
    f=filtros or {}; cond,p=_filtros_relatorio(f); cond.append("r.status='REJEITADA'")
    if f.get("motivo"): cond.append("LOWER(r.motivo_rejeicao) LIKE ?"); p.append("%"+str(f["motivo"]).lower()+"%")
    conn=conectar()
    try:
        cur=conn.cursor(); cur.execute(q(f"""SELECT r.*,(SELECT e.usuario_nome FROM requisicao_compra_eventos e WHERE e.requisicao_compra_id=r.id AND e.evento='REJEICAO' ORDER BY e.id DESC LIMIT 1) decisor_nome FROM requisicoes_compra r WHERE {' AND '.join(cond)} ORDER BY r.id DESC"""),tuple(p)); linhas=[]
        for row in cur.fetchall(): rc=dict(row); cur.execute(q("SELECT * FROM requisicao_compra_itens WHERE requisicao_compra_id=?"),(rc["id"],)); rc["itens"]=[dict(x) for x in cur.fetchall()]; from .services import _totais; _totais(rc); rc["gasto_evitado"]=rc["total_estimado_solicitado"]; linhas.append(rc)
        total=sum((x["gasto_evitado"] for x in linhas),Decimal("0")); setores={}; motivos={}
        for x in linhas: setores[x["setor"]]=setores.get(x["setor"],Decimal("0"))+x["gasto_evitado"]; motivos[x.get("motivo_rejeicao") or "Não informado"]=motivos.get(x.get("motivo_rejeicao") or "Não informado",0)+1
        return {"linhas":linhas,"indicadores":{"quantidade":len(linhas),"total":total,"ticket":total/len(linhas) if linhas else Decimal("0"),"setor":max(setores,key=setores.get) if setores else "—","motivo":max(motivos,key=motivos.get) if motivos else "—"}}
    finally: conn.close()

def gerar_excel(tipo,filtros=None):
    dados=relatorio_atendidas(filtros) if tipo=="atendidas" else relatorio_recusadas(filtros); wb=Workbook(); ws=wb.active; ws.title="RCs "+tipo.title()
    if tipo=="atendidas": cab=("RC","Solicitação","Solicitante","Setor","Fornecedor","Aprovação","Atendimento","Estimado","Real","Diferença","Variação %","Documentos","Status"); rows=[(x["numero"],x["criado_em"],x["solicitante_nome_snapshot"],x["setor"],x["fornecedores"],x["aprovado_em"],x["data_atendimento"],float(x["valor_estimado"]),float(x["valor_real"]),float(x["diferenca"]),None if x["variacao"] is None else float(x["variacao"]),x["documentos"],x["status"]) for x in dados["linhas"]]
    else: cab=("RC","Data","Solicitante","Setor","Motivo","Gasto evitado","Decisor","Observação"); rows=[(x["numero"],x["criado_em"],x["solicitante_nome_snapshot"],x["setor"],x["motivo_rejeicao"],float(x["gasto_evitado"]),x.get("decisor_nome") or "—",x.get("observacoes")) for x in dados["linhas"]]
    ws.append(cab)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1267B2")
    for row in rows: ws.append(row)
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(38,max(12,max(len(str(c.value or "")) for c in col)+2))
    out=BytesIO(); wb.save(out); return out.getvalue()

def resumo_portal():
    atend=relatorio_atendidas({}); recus=relatorio_recusadas({}); conn=conectar()
    try:
        cur=conn.cursor(); cur.execute("SELECT status,COUNT(*) total FROM requisicoes_compra GROUP BY status"); s={x["status"]:x["total"] for x in cur.fetchall()}; return {"aguardando":s.get("APROVADA",0)+s.get("APROVADA_COM_AJUSTES",0),"parciais":s.get("PARCIALMENTE_ATENDIDA",0),"atendidas":s.get("ATENDIDA",0),"estimado":atend["indicadores"]["estimado"],"real":atend["indicadores"]["real"],"gasto_evitado":recus["indicadores"]["total"]}
    finally: conn.close()
